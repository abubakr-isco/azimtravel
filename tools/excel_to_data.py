"""
Импорт данных из Excel-файла «Графики июль–декабрь 2026» в js/data.js для демо CRM.

Запуск (нужен openpyxl):
    pip install openpyxl
    python3 tools/excel_to_data.py "Графики_июль_декабрь_2026_г —_копия.xlsx"
"""
import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "js" / "data.js"

# кириллические буквы, похожие на латинские, в госномерах
CYR_TO_LAT = str.maketrans("АВСЕКМНОРТХУ", "ABCEKMHOPTXY")

# месяц узнаём по первым трём буквам: в шапке бывают опечатки («августь »)
MONTHS = {"янв": 1, "фев": 2, "мар": 3, "апр": 4, "май": 5, "мая": 5, "июн": 6,
          "июл": 7, "авг": 8, "сен": 9, "окт": 10, "ноя": 11, "дек": 12}


def month_of(v):
    return MONTHS.get(v.strip().lower()[:3]) if isinstance(v, str) else None


def clean(v):
    """Пустые значения и нули-заглушки из формул -> None, строки -> обрезанные."""
    if v is None:
        return None
    if isinstance(v, str):
        v = re.sub(r"\s+", " ", v).strip()
        return v or None
    if isinstance(v, (int, float)) and v == 0:
        return None
    return v


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return None


def as_time(v):
    if isinstance(v, dt.time):
        return v.strftime("%H:%M")
    if isinstance(v, dt.datetime):
        return v.strftime("%H:%M")
    if isinstance(v, str) and re.fullmatch(r"\d{1,2}[:^.]\d{2}", v.strip()):
        return re.sub(r"[\^.]", ":", v.strip()).zfill(5)
    return None


def as_num(v):
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    if isinstance(v, str):
        try:
            return round(float(v.replace(" ", "").replace(",", ".")), 2)
        except ValueError:
            return None
    return None


def services(ws, rows, left_col, right_col):
    """Услуги дня: в левой колонке время, в правой — услуга (вс аэ, пр ж/в, экя...)."""
    out = []
    for r in rows:
        left, right = ws.cell(r, left_col).value, ws.cell(r, right_col).value
        t = as_time(left)
        name = clean(right) if isinstance(right, str) else None
        if left is not None and t is None and isinstance(clean(left), str):
            # иногда услуга записана в левой колонке без времени
            out.append({"t": None, "s": clean(left)})
        if name:
            out.append({"t": t, "s": name})
        elif t:
            out.append({"t": t, "s": None})
    return out


def is_cancel(*vals):
    return any(isinstance(v, str) and "аннул" in v.lower() for v in vals)


def read_requests(wb):
    ws = wb["Заявки"]
    pay = wb["Взаиморасчеты (заказч)"]
    requests = []
    for k, main in enumerate(range(6, ws.max_row + 1, 6)):
        customer = clean(ws.cell(main, 2).value)
        if not customer or customer == "0":
            continue
        no = ws.cell(main, 1).value
        date_from = as_date(ws.cell(main, 5).value)
        date_to = as_date(ws.cell(main, 6).value)
        if not date_from or date_from < "2026-01-01":
            print(f"  пропущена заявка №{no} ({customer}): дата {date_from}")
            continue
        country = clean(ws.cell(main, 7).value)
        escort = clean(ws.cell(main, 9).value)
        cancelled = is_cancel(country, escort)
        if is_cancel(country):
            country = None
        if is_cancel(escort):
            escort = None

        days = []
        start = dt.date.fromisoformat(date_from) if date_from else None
        for i in range(14):
            lc, rc = 10 + 2 * i, 11 + 2 * i
            route = clean(ws.cell(main, lc).value)
            svc = services(ws, range(main + 1, main + 4), lc, rc)
            cost = as_num(ws.cell(main + 4, lc).value)
            if not (route or svc or cost):
                continue
            d = as_date(ws.cell(main - 1, lc).value)
            if not d and start:
                d = (start + dt.timedelta(days=i)).isoformat()
            days.append({"n": i + 1, "d": d, "route": route if isinstance(route, str) else None,
                         "svc": svc, "cost": cost or 0})

        total = as_num(ws.cell(main + 4, 38).value)
        if total is None:
            total = sum(d["cost"] for d in days)

        # строка взаиморасчётов с тем же порядковым номером
        pr = 6 + k
        payments = []
        rate = as_num(pay.cell(pr, 11).value) or 12900
        vat = as_num(pay.cell(pr, 12).value) or 12
        cash = as_num(pay.cell(pr, 15).value)
        if cash:
            payments.append({"date": as_date(pay.cell(pr, 14).value), "method": "cash",
                             "usd": cash, "note": "Наличные USD"})
        uzs = as_num(pay.cell(pr, 18).value)
        if uzs:
            p_rate = as_num(pay.cell(pr, 19).value) or rate
            p_vat = as_num(pay.cell(pr, 20).value) or vat
            annex = clean(pay.cell(pr, 17).value)
            payments.append({"date": None, "method": "transfer", "uzs": uzs, "rate": p_rate, "vat": p_vat,
                             "usd": round(uzs / p_rate / (1 + p_vat / 100), 2),
                             "note": f"Перечисление, приложение №{annex}" if annex else "Перечисление"})

        transport = clean(ws.cell(main, 3).value)
        pax = as_num(ws.cell(main, 4).value)
        group = clean(ws.cell(main, 8).value)
        requests.append({
            "id": f"R{k + 1}",
            "no": int(no) if isinstance(no, (int, float)) else k + 1,
            "customer": customer,
            "transport": str(transport) if transport is not None else None,
            "pax": int(pax) if pax else None,
            "from": date_from,
            "to": date_to,
            "country": country,
            "group": str(group) if group is not None else None,
            "escort": escort,
            "cancelled": cancelled,
            "total": 0 if cancelled and not total else total,
            "rate": rate,
            "vat": vat,
            "days": days,
            "payments": payments,
        })
    return requests


def parse_vehicle(name):
    brand = name.split()[0].upper() if name.split()[0].isupper() and name.split()[0].isascii() else None
    plate = re.search(r"(\d{2}\s+)?\d{3}\s*[A-ZА-Я]{3}", name)
    rest = name
    if brand:
        rest = rest[len(name.split()[0]):]
    if plate:
        rest = rest.replace(plate.group(0), "")
    seats = re.search(r"\b(55\+2|VIP)\b", rest)
    if seats:
        rest = rest.replace(seats.group(0), "")
    sold = bool(re.search(r"продан", name, re.I))
    rest = re.sub(r"продан[оа]", "", rest, flags=re.I)
    rest = re.sub(r"/\s*11\s*[MМ]", "", rest)
    driver = re.sub(r"\s+", " ", rest).strip(" /") or None
    return {
        "brand": brand or "Партнёр",
        "plate": re.sub(r"\s+", " ", plate.group(0)).translate(CYR_TO_LAT) if plate else None,
        "driver": driver,
        "seats": seats.group(0) if seats else None,
        "sold": sold,
        "partner": brand is None,
    }


def read_vehicles(wb):
    ws = wb["График транспортный"]
    st = wb["Взаиморасчеты (трансп)"]
    # колонка -> дата по шапке (строка 3 — месяц, строка 4 — число)
    col_date = {}
    month = None
    for c in range(3, ws.max_column + 1, 2):
        m = ws.cell(3, c).value
        if month_of(m):
            month = month_of(m)
        elif isinstance(m, str) and m.strip():
            print(f"  непонятный месяц в шапке графика транспорта: {m!r} (колонка {c})")
        day = ws.cell(4, c).value
        if month and isinstance(day, (int, float)):
            d = dt.date(2026, month, int(day)).isoformat()
            if d in col_date.values():
                print(f"  дата {d} встречается в шапке графика транспорта дважды (колонка {c})")
            col_date[c] = d

    paid_by_no = {}
    for r in range(7, st.max_row + 1):
        no = st.cell(r, 1).value
        if isinstance(no, (int, float)):
            paid_by_no[int(no)] = {"cost": as_num(st.cell(r, 3).value) or 0,
                                   "paid": as_num(st.cell(r, 5).value) or 0,
                                   "date": as_date(st.cell(r, 4).value)}

    vehicles = []
    for top in range(5, ws.max_row + 1, 5):
        name = clean(ws.cell(top, 2).value)
        no = ws.cell(top, 1).value
        if not name:
            continue
        schedule = []
        for c, d in col_date.items():
            route = clean(ws.cell(top, c).value)
            svc = services(ws, range(top + 1, top + 4), c, c + 1)
            cost = as_num(ws.cell(top + 4, c).value)
            if not (route or svc or cost):
                continue
            schedule.append({"d": d, "route": route if isinstance(route, str) else None,
                             "svc": svc, "cost": cost or 0})
        info = parse_vehicle(name)
        s = paid_by_no.get(int(no), {}) if isinstance(no, (int, float)) else {}
        payments = []
        if s.get("paid"):
            payments.append({"date": s.get("date"), "usd": s["paid"], "note": "Оплата перевозчику"})
        vehicles.append({
            "id": f"V{int(no)}",
            "no": int(no),
            "name": re.sub(r"\s+", " ", name),
            **info,
            "schedule": schedule,
            "cost": s.get("cost", sum(x["cost"] for x in schedule)),
            "payments": payments,
        })
    return vehicles


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else next(ROOT.glob("*.xlsx"))
    wb = openpyxl.load_workbook(src, data_only=True)
    data = {
        "source": src.name,
        "importedAt": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "settings": {"rate": 12900, "vat": 12, "company": "Azim Travel"},
        "requests": read_requests(wb),
        "vehicles": read_vehicles(wb),
    }
    OUT.write_text("// Сгенерировано tools/excel_to_data.py — не редактировать вручную\n"
                   "window.AZIM_DATA = " + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n",
                   encoding="utf-8")
    print(f"Заявок: {len(data['requests'])}, транспорт: {len(data['vehicles'])} -> {OUT}")


if __name__ == "__main__":
    main()
